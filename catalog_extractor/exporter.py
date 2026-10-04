"""Write product variants to a styled .xlsx file."""

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import COLUMN_HEADERS, COLUMN_ORDER, ProductVariant

logger = logging.getLogger(__name__)

COLUMN_WIDTHS = {
    "section": 16, "category": 26, "product_line": 34, "product_name": 48,
    "ean": 16, "code": 16, "size": 18, "packaging": 14, "price": 10,
}


def export_xlsx(products: list[ProductVariant], output_dir: str = "output") -> str:
    """Write `products` to `<output_dir>/catalog_<timestamp>.xlsx` and return the path."""
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / f"catalog_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Catalog"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, key in enumerate(COLUMN_ORDER, start=1):
        cell = ws.cell(row=1, column=col, value=COLUMN_HEADERS[key])
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = COLUMN_WIDTHS[key]
    ws.freeze_panes = "A2"

    for row, product in enumerate(products, start=2):
        for col, key in enumerate(COLUMN_ORDER, start=1):
            value = getattr(product, key)
            cell = ws.cell(row=row, column=col, value=value)
            if key == "price" and value is not None:
                cell.number_format = "#,##0.00"
            elif key in ("ean", "code"):
                cell.number_format = "@"  # keep leading zeros
    ws.auto_filter.ref = ws.dimensions

    wb.save(out_file)
    logger.info("exported %d rows to %s", len(products), out_file)
    return str(out_file)
