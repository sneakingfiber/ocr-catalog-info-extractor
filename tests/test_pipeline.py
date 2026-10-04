from pathlib import Path

from openpyxl import load_workbook

from catalog_extractor import Profile, classify_record, export_xlsx, normalize, parse_ocr_text
from catalog_extractor.parser import parse_price

ROOT = Path(__file__).parent.parent
SAMPLE = (ROOT / "examples" / "sample_ocr_page.txt").read_text()


def test_format_a_inline_ean():
    recs = parse_ocr_text(SAMPLE)
    r = recs[0]
    assert r["ean"] == "8000000000011"
    assert r["code"] == "1000001"
    assert r["size"] == "16cm"
    assert r["packaging"] == "pz. 6"
    assert r["price"] == 16.61
    assert "ROUND TIP" in r["product_name"]


def test_format_b_ean_on_following_line():
    recs = parse_ocr_text(SAMPLE)
    level = [r for r in recs if r["code"] in ("323", "324")]
    assert [r["ean"] for r in level] == ["8000000000035", "8000000000042"]


def test_code_without_ean_is_unchanged():
    recs = parse_ocr_text(SAMPLE)
    r = [r for r in recs if r["code"] == "40622MF"][0]
    assert r["ean"] is None


def test_record_count():
    assert len(parse_ocr_text(SAMPLE)) == 5


def test_price_formats():
    assert parse_price("16,61") == 16.61
    assert parse_price("1.204,00") == 1204.0
    assert parse_price("16.61") == 16.61
    assert parse_price("abc") is None


def test_classification():
    profile = Profile.load(str(ROOT / "profiles" / "example.yaml"))
    rec = classify_record(parse_ocr_text(SAMPLE)[0], page=7, profile=profile)
    assert rec["section"] == "BUILDING"
    assert rec["category"] == "MASON'S TOOLS"
    assert rec["product_line"] == 'MASON\'S TOOLS - LINE "ANCHOR" ACME'
    assert classify_record(parse_ocr_text(SAMPLE)[0], 100, profile)["section"] == "SITE EQUIPMENT"


def test_empty_profile_does_not_crash():
    rec = classify_record(parse_ocr_text(SAMPLE)[0], page=1, profile=Profile())
    assert rec["category"] == "TROWEL"
    assert rec["section"] == ""


def test_normalize_dedup_and_code_preserved():
    recs = parse_ocr_text(SAMPLE)
    out = normalize(recs + recs)
    assert len(out) == 5
    assert "40622MF" in [v.code for v in out]


def test_export_roundtrip(tmp_path):
    out = normalize(parse_ocr_text(SAMPLE))
    path = export_xlsx(out, str(tmp_path))
    ws = load_workbook(path).active
    assert ws.max_row == 6
    assert ws["A1"].value == "SECTION"
    assert ws["I2"].value == 16.61
